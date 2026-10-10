"""
Bước 5: gộp bài phân tích thành file Word.

    .\\venv\\Scripts\\python.exe xuat_word.py
    .\\venv\\Scripts\\python.exe xuat_word.py --laso laso\\ten_file.json

File ra nằm trong thư mục output\\năm-tháng-ngày_giờsinh\\TuVi_Report.docx
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from models import load_chart, thu_muc_bao_cao

ROOT = Path(__file__).resolve().parent
MAU_IN = re.compile(r"\*\*(.+?)\*\*|\*(.+?)\*")


def gan_font(run, dam: bool = False, nghieng: bool = False) -> None:
    run.bold = dam
    run.italic = nghieng
    run.font.name = "Times New Roman"
    run.font.size = Pt(13)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")


def mo_link(van: str) -> str:
    """Đổi link trong bài thành chữ và địa chỉ trang, để file Word bấm được bằng mắt."""
    return re.sub(
        r"\[([^\]]+)\]\((/[^)]+)\)",
        lambda khop: f"{khop.group(1)} (http://127.0.0.1:8765{khop.group(2)})",
        van,
    )


def them_chu(doan, van: str, dam: bool = False) -> None:
    van = mo_link(van)
    vi_tri = 0
    for khop in MAU_IN.finditer(van):
        if khop.start() > vi_tri:
            gan_font(doan.add_run(van[vi_tri : khop.start()]), dam)
        if khop.group(1) is not None:
            gan_font(doan.add_run(khop.group(1)), True)
        else:
            gan_font(doan.add_run(khop.group(2)), dam, True)
        vi_tri = khop.end()
    if vi_tri < len(van):
        gan_font(doan.add_run(van[vi_tri:]), dam)


def la_dong_bang(dong: str) -> bool:
    return dong.startswith("|") and dong.endswith("|") and dong.count("|") >= 3


def la_dong_ke(dong: str) -> bool:
    ruot = dong.replace("|", "").replace("-", "").replace(":", "").strip()
    return ruot == ""


def o_bang(dong: str) -> list[str]:
    return [o.strip() for o in dong.strip("|").split("|")]


def gop_mot_hang(dong: list[str], i: int) -> tuple[str, int]:
    """Một hàng bảng đôi khi bị xuống dòng giữa ô. Gom đến khi gặp dấu | đóng hàng."""
    gop = dong[i].strip()
    i += 1
    while not gop.endswith("|") and i < len(dong):
        tiep = dong[i].strip()
        if tiep.startswith("|") and tiep.endswith("|") and tiep.count("|") >= 3:
            break
        if tiep.startswith("## "):
            break
        i += 1
        if not tiep:
            continue
        gop += "<br>" + tiep
    return gop, i


def ve_bang(tai_lieu, hang: list[list[str]]) -> None:
    if not hang:
        return
    so_cot = max(len(mot) for mot in hang)
    bang = tai_lieu.add_table(rows=len(hang), cols=so_cot)
    bang.style = "Table Grid"
    for i, mot in enumerate(hang):
        for j in range(so_cot):
            chu = mot[j] if j < len(mot) else ""
            o = bang.cell(i, j)
            o.text = ""
            doan = o.paragraphs[0]
            phan = re.split(r"\s*<br\s*/?>\s*", chu, flags=re.I)
            for k, mot_phan in enumerate(phan):
                if k:
                    ngat = doan.add_run()
                    gan_font(ngat, i == 0)
                    ngat.add_break()
                them_chu(doan, mot_phan, i == 0)


def them_thong_tin(tai_lieu, laso: Path | None) -> None:
    chart = load_chart(laso)
    tieu_de = tai_lieu.add_heading("Báo cáo lá số Tử Vi", level=0)
    for run in tieu_de.runs:
        gan_font(run, True)
        run.font.size = Pt(18)
    dong = [
        f"Giới tính: {chart.gioi_tinh}. {chart.am_duong}.",
        f"Sinh: {chart.ngay_sinh}, giờ hành chính {chart.gio_hanh_chinh}, giờ mặt trời {chart.gio_mat_troi}.",
        f"Âm lịch: {chart.am_lich}. Nơi sinh: {chart.noi_sinh}.",
        f"Bản mệnh {chart.ban_menh}. Cục {chart.cuc}, khởi vận {chart.khoi_van_tuoi} tuổi. Thân cư {chart.than_cu}.",
        "Phần dưới chỉ tổng hợp từ sách Tử Vi Đẩu Số Tân Biên.",
    ]
    for cau in dong:
        doan = tai_lieu.add_paragraph()
        them_chu(doan, cau)


def xuat(laso: str | None = None) -> None:
    chart_path = None
    if laso:
        chart_path = Path(laso)
        if not chart_path.is_file():
            chart_path = ROOT / "laso" / laso
        if not chart_path.is_file():
            raise SystemExit(f"Không thấy file lá số: {laso}")
        chart = load_chart(chart_path)
    else:
        chart = load_chart()
    thu_muc = thu_muc_bao_cao(chart)
    nguon = thu_muc / "phan_tich.md"
    dich = thu_muc / "TuVi_Report.docx"
    if not nguon.is_file():
        raise SystemExit(f"Chưa có bài phân tích: {nguon}")
    tai_lieu = Document()
    for section in tai_lieu.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)
    them_thong_tin(tai_lieu, chart_path)

    cac_dong = nguon.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(cac_dong):
        dong = cac_dong[i].strip()
        if not dong or dong == "# Phân tích lá số" or dong.startswith("Cách viết:"):
            i += 1
            continue
        if dong.startswith("|"):
            hang = []
            while i < len(cac_dong) and cac_dong[i].strip().startswith("|"):
                gop, i = gop_mot_hang(cac_dong, i)
                if not la_dong_ke(gop.replace("<br>", "")):
                    hang.append(o_bang(gop))
            ve_bang(tai_lieu, hang)
            continue
        if dong.startswith("## "):
            tieu = tai_lieu.add_heading(dong[3:].strip(), level=1)
            for run in tieu.runs:
                gan_font(run, True)
                run.font.size = Pt(16)
            i += 1
            continue
        if dong.startswith("* ") or dong.startswith("- "):
            doan = tai_lieu.add_paragraph(style="List Bullet")
            them_chu(doan, dong[2:].strip())
            i += 1
            continue
        doan = tai_lieu.add_paragraph()
        them_chu(doan, dong)
        i += 1

    tai_lieu.save(dich)
    print(f"Xong. Mở file: {dich}")


if __name__ == "__main__":
    import argparse

    tham_so = argparse.ArgumentParser(description="Xuất bài phân tích ra file Word.")
    tham_so.add_argument("--laso", help="Cùng file lá số đã dùng với analyze.py, nếu không phải lá số mẫu.")
    xuat(tham_so.parse_args().laso)
